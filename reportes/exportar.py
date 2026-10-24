"""
Exportaciones de reportes: CSV, Excel (openpyxl) y PDF (reportlab).

Equivalen a AdminController.exportarCsv / exportarPdfCompleto / exportarPdfCliente
y a ReporteController.exportarExcel. El PDF reproduce el diseño de iText (tamaños,
paleta, encabezado con logo, línea verde, tablas con filas alternas y estado en
color); es equivalente en contenido, no idéntico píxel a píxel.
"""
import csv
from io import BytesIO, StringIO
from pathlib import Path
from xml.sax.saxutils import escape

from django.conf import settings
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core.models import Cliente, Sensor, Vehiculo
from core.patrones.fabrica import EstadoSensorStrategyFactory, TipoEstado
from core.utilidades import ahora, texto_java
from reportes.filtros import (filtrar_clientes, filtrar_sensores, limitar_a_clientes, normalizar_rango, nz_nivel,
                              vehiculos_con_falla, vehiculos_de)

STRATEGY_REPORTE = EstadoSensorStrategyFactory.crear(TipoEstado.REPORTE)

# Paleta AutoSen (AdminController.java)
NEGRO = colors.Color(13 / 255, 13 / 255, 15 / 255)
AZUL_OSCURO = colors.Color(15 / 255, 32 / 255, 61 / 255)
AZUL_MEDIO = colors.Color(41 / 255, 98 / 255, 168 / 255)
BLANCO = colors.white
VERDE = colors.Color(46 / 255, 184 / 255, 120 / 255)
GRIS_CLARO = colors.Color(240 / 255, 243 / 255, 247 / 255)
GRIS_TEXTO = colors.Color(90 / 255, 96 / 255, 105 / 255)

LOGO = Path(settings.BASE_DIR) / "static" / "images" / "logo-autosen.png"
MARGEN = 30
SALTO_LINEA = 12  # equivale a doc.add(new Paragraph("\n"))
FIN_LINEA = "\r\n"  # println/%n de Java en Windows


def _nr(texto) -> str:
    return "" if texto is None else str(texto)


def _nombre_completo(cliente) -> str:
    return f"{_nr(cliente.nombre_cliente)} {_nr(cliente.apellido_cliente)}"


# ---------------------------------------------------------------------------
# CSV (AdminController.exportarCsv)
# ---------------------------------------------------------------------------
def csv_reporte(tipo: str) -> str:
    """Mismas columnas que los printf de Java, pero escapado con el módulo csv, con vacío en lugar de "null"
    y con BOM para que Excel lea las tildes (lote A). Un tipo desconocido da un archivo vacío, como en Java."""
    salida = StringIO()
    fila = csv.writer(salida, lineterminator=FIN_LINEA).writerow
    if tipo == "clientes":
        fila(["ID", "Nombre", "Correo", "Estado", "Vehiculos"])
        for c in Cliente.objects.all():
            fila([c.id_cliente, _nombre_completo(c), _nr(c.correo), _nr(c.estado), c.vehiculos.count()])
    elif tipo == "sensores":
        fila(["Sensor", "Tipo", "Nivel", "Estado", "Vehiculo", "Cliente"])
        for s in Sensor.objects.all():
            estado = "OK" if s.nivel < 40 else "Advertencia" if s.nivel < 70 else "Falla"
            vehiculo = _nr(s.vehiculo.nombre_vehiculo) if s.vehiculo else "-"
            cliente = _nr(s.vehiculo.cliente.nombre_cliente) if s.vehiculo and s.vehiculo.cliente else "-"
            fila([_nr(s.nombre_sensor), _nr(s.tipo_sensor), s.nivel, estado, vehiculo, cliente])
    elif tipo == "fallas":
        fila(["Vehiculo", "Placa", "Cliente", "Sensor", "Nivel", "Tipo Dano"])
        for s in Sensor.objects.all():
            if s.nivel >= 70:
                vehiculo = _nr(s.vehiculo.nombre_vehiculo) if s.vehiculo else "-"
                placa = _nr(s.vehiculo.placa) if s.vehiculo else "-"
                cliente = _nr(s.vehiculo.cliente.nombre_cliente) if s.vehiculo and s.vehiculo.cliente else "-"
                tipo_dano = s.tipo_dano if s.tipo_dano is not None else "-"
                fila([vehiculo, placa, cliente, _nr(s.nombre_sensor), s.nivel, tipo_dano])
    elif tipo == "bloqueados":
        fila(["Nombre", "Correo", "Vehiculos"])
        for c in Cliente.objects.all():
            if c.estado == "bloqueado":
                fila([_nombre_completo(c), _nr(c.correo), c.vehiculos.count()])
    texto = salida.getvalue()
    return "\ufeff" + texto if texto else ""


# ---------------------------------------------------------------------------
# Excel (ReporteController.exportarExcel)
# ---------------------------------------------------------------------------
def _fila(hoja, *valores) -> None:
    """crearFila de Java: null se escribe vacío; números como números."""
    hoja.append(["" if v is None else v for v in valores])


def excel_reporte(clientes, sensores) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Clientes"
    _fila(hoja, "ID", "Nombre", "Correo", "Estado", "Vehículos")
    for c in clientes:
        _fila(hoja, c.id_cliente, _nombre_completo(c), _nr(c.correo), _nr(c.estado), c.vehiculos.count())

    hoja = libro.create_sheet("Sensores")
    _fila(hoja, "Sensor", "Tipo", "Vehículo", "Cliente", "Nivel", "Estado", "Diagnóstico")
    for s in sensores:
        v = s.vehiculo
        cliente = _nombre_completo(v.cliente) if v is not None and v.cliente is not None else "-"
        _fila(hoja, _nr(s.nombre_sensor), _nr(s.tipo_sensor), _nr(v.nombre_vehiculo) if v else "-", cliente,
              nz_nivel(s), STRATEGY_REPORTE.clasificar(s.nivel), _nr(s.tipo_dano))

    # Las dos últimas hojas cuentan lo mismo que /admin/reportes-filtrado (revisión final del lote A).
    hoja = libro.create_sheet("Vehículos con fallas")
    _fila(hoja, "Vehículo", "Placa", "Cliente")
    for v in vehiculos_con_falla(vehiculos_de(Vehiculo.objects.all(), clientes), sensores):
        _fila(hoja, _nr(v.nombre_vehiculo), _nr(v.placa), _nombre_completo(v.cliente))

    hoja = libro.create_sheet("Clientes bloqueados")
    _fila(hoja, "Nombre", "Correo", "Vehículos")
    for c in clientes:
        if c.estado == "bloqueado":
            _fila(hoja, _nombre_completo(c), _nr(c.correo), c.vehiculos.count())

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# PDF (AdminController.exportarPdfCompleto / exportarPdfCliente)
# ---------------------------------------------------------------------------
def datos_pdf_general(params: dict) -> tuple[list, list]:
    """Clientes y sensores del PDF general con los mismos filtros que la vista y el Excel (F5)."""
    nivel_min, nivel_max = normalizar_rango(params.get("nivelMin"), params.get("nivelMax"))
    clientes = filtrar_clientes(list(Cliente.objects.all()), params.get("q"), params.get("estadoCliente"),
                                params.get("marca"), params.get("tipoPlaca"))
    sensores = filtrar_sensores(list(Sensor.objects.all()), params.get("estadoSensor"), params.get("tipoSensor"),
                                nivel_min, nivel_max)
    return clientes, limitar_a_clientes(sensores, clientes, params)


def _estilo(nombre: str, fuente: str, tamano: float, color) -> ParagraphStyle:
    return ParagraphStyle(nombre, fontName=fuente, fontSize=tamano, leading=tamano * 1.25, textColor=color)


def _parrafo(texto, fuente="Helvetica", tamano=9, color=NEGRO) -> Paragraph:
    return Paragraph(escape(_nr(texto)), _estilo("celda", fuente, tamano, color))


def _color_estado(estado: str):
    return {"OK": VERDE, "ADVERTENCIA": AZUL_MEDIO}.get(estado, NEGRO)


def _encabezado(titulo: str, ancho: float) -> list:
    """Título a la izquierda, logo (si existe) a la derecha y línea verde debajo (agregarEncabezado)."""
    logo = Image(str(LOGO), width=60, height=60, kind="proportional") if LOGO.is_file() else ""
    cabecera = Table([[_parrafo(titulo, "Helvetica-Bold", 20, AZUL_OSCURO), logo]],
                     colWidths=[ancho * 3 / 4, ancho / 4])
    cabecera.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (1, 0), (1, 0), "RIGHT")]))
    linea = Table([[""]], colWidths=[ancho], rowHeights=[3])
    linea.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), VERDE)]))
    return [cabecera, linea, Spacer(1, 10)]


def _fecha() -> list:
    """agregarFecha: "Generado el dd/MM/yyyy HH:mm" y una línea en blanco."""
    return [_parrafo("Generado el " + ahora().strftime("%d/%m/%Y %H:%M"), tamano=9, color=GRIS_TEXTO),
            Spacer(1, SALTO_LINEA)]


def _tabla(columnas: list[str], filas: list[list], ancho: float, columna_estado: int | None,
           color_estado=_color_estado) -> Table:
    """Cabecera azul oscuro y filas alternas blanco / gris claro; el estado va en negrita y en color."""
    datos = [[_parrafo(c, "Helvetica-Bold", 9, BLANCO) for c in columnas]]
    estilos = [
        ("BACKGROUND", (0, 0), (-1, 0), AZUL_OSCURO),
        ("GRID", (0, 0), (-1, 0), 0.5, NEGRO),
        ("GRID", (0, 1), (-1, -1), 0.5, GRIS_CLARO),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        *[(lado, (0, 0), (-1, 0), 7) for lado in ("LEFTPADDING", "RIGHTPADDING", "TOPPADDING", "BOTTOMPADDING")],
        *[(lado, (0, 1), (-1, -1), 6) for lado in ("LEFTPADDING", "RIGHTPADDING", "TOPPADDING", "BOTTOMPADDING")],
    ]
    for i, fila in enumerate(filas, start=1):
        celdas = []
        for j, valor in enumerate(fila):
            if j == columna_estado:
                celdas.append(_parrafo(valor, "Helvetica-Bold", 9, color_estado(valor)))
            else:
                celdas.append(_parrafo(valor))
        datos.append(celdas)
        estilos.append(("BACKGROUND", (0, i), (-1, i), BLANCO if i % 2 == 1 else GRIS_CLARO))
    tabla = Table(datos, colWidths=[ancho / len(columnas)] * len(columnas), repeatRows=1)
    tabla.setStyle(TableStyle(estilos))
    return tabla


def _documento(tamano) -> tuple[BytesIO, SimpleDocTemplate]:
    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=tamano, leftMargin=MARGEN, rightMargin=MARGEN, topMargin=MARGEN,
                            bottomMargin=MARGEN)
    return salida, doc


def pdf_general(clientes, sensores) -> bytes:
    """REPORTE GENERAL AUTOSEN: tabla de clientes y estado de sensores, A4 horizontal."""
    salida, doc = _documento(landscape(A4))
    ancho = doc.width
    estilo_seccion = _estilo("seccion", "Helvetica-Bold", 13, AZUL_OSCURO)

    filas_clientes = [[c.id_cliente, _nombre_completo(c), c.correo, c.estado, c.vehiculos.count()] for c in clientes]
    # En la tabla de clientes el estado va en verde si está activo y en negro si no.
    tabla_clientes = _tabla(["ID", "Nombre", "Correo", "Estado", "Vehículos"], filas_clientes, ancho, 3,
                            color_estado=lambda estado: VERDE if estado == "activo" else NEGRO)

    filas_sensores = []
    for s in sensores:
        v = s.vehiculo
        filas_sensores.append([
            s.nombre_sensor, s.tipo_sensor,
            v.nombre_vehiculo if v else "-",
            v.cliente.nombre_cliente if v and v.cliente else "-",
            f"{nz_nivel(s)}%", STRATEGY_REPORTE.clasificar(nz_nivel(s)),
        ])

    elementos = [
        *_encabezado("REPORTE GENERAL AUTOSEN", ancho), *_fecha(),
        Paragraph("CLIENTES", estilo_seccion),
        Spacer(1, SALTO_LINEA), tabla_clientes, Spacer(1, SALTO_LINEA),
        Paragraph("ESTADO DE SENSORES", estilo_seccion),
        Spacer(1, SALTO_LINEA),
        _tabla(["Sensor", "Tipo", "Vehículo", "Cliente", "Nivel", "Estado"], filas_sensores, ancho, 5),
    ]
    doc.build(elementos)
    return salida.getvalue()


def pdf_cliente(cliente) -> bytes:
    """Reporte de un cliente: un bloque por vehículo con la tabla de sus sensores, A4 vertical."""
    salida, doc = _documento(A4)
    ancho = doc.width
    elementos = [*_encabezado(f"{_nr(cliente.nombre_cliente)} {_nr(cliente.apellido_cliente)}", ancho), *_fecha()]
    for v in cliente.vehiculos.all():
        elementos.append(Spacer(1, SALTO_LINEA))
        elementos.append(_parrafo(f"{texto_java(v.nombre_vehiculo)} — {texto_java(v.marca)}",
                                  "Helvetica-Bold", 13, AZUL_OSCURO))
        elementos.append(_parrafo(f"Placa: {v.placa if v.placa is not None else '-'}   |   Color: {texto_java(v.color)}",
                                  tamano=10, color=GRIS_TEXTO))
        sensores = list(v.sensores.all())
        if sensores:
            filas = []
            for s in sensores:
                nivel = nz_nivel(s)
                filas.append([s.nombre_sensor, s.tipo_sensor, f"{nivel}%", STRATEGY_REPORTE.clasificar(nivel),
                              s.tipo_dano if s.tipo_dano is not None else "-"])
            elementos += [Spacer(1, 8), _tabla(["Sensor", "Tipo", "Nivel", "Estado", "Tipo Daño"], filas, ancho, 3),
                          Spacer(1, 10)]
    doc.build(elementos)
    return salida.getvalue()
