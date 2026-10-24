import csv
import io
from io import BytesIO
from unittest import mock

from django.test import TestCase
from openpyxl import load_workbook

from core.models import Cliente
from core.tests.fabricas import crear_cliente, crear_usuario, crear_vehiculo, login_admin
from reportes import exportar
from reportes.exportar import datos_pdf_general
from reportes.tests.test_filtros import datos_reportes


class ExportarTests(TestCase):
    def setUp(self):
        login_admin(self.client, crear_usuario())
        self.joel, self.ana, (self.s1, self.s2, self.s3, self.s4) = datos_reportes()

    def test_csv_clientes(self):
        r = self.client.get("/admin/reportes/csv/clientes")
        self.assertEqual((r["Content-Type"], r["Content-Disposition"]),
                         ("text/csv; charset=UTF-8", "attachment; filename=reporte-clientes.csv"))
        self.assertEqual(r.content.decode(),
                         f"\ufeffID,Nombre,Correo,Estado,Vehiculos\r\n{self.joel.pk},Joel Moncaleano,joel@test.com,activo,1\r\n"
                         f"{self.ana.pk},Ana Paz,ana@test.com,bloqueado,1\r\n")

    def test_csv_fallas_bloqueados_y_desconocido(self):
        self.assertEqual(self.client.get("/admin/reportes/csv/fallas").content.decode(),
                         "\ufeffVehiculo,Placa,Cliente,Sensor,Nivel,Tipo Dano\r\nRio,KIA111,Ana,Sensor ECT,80,-\r\n")
        self.assertEqual(self.client.get("/admin/reportes/csv/bloqueados").content.decode(),
                         "\ufeffNombre,Correo,Vehiculos\r\nAna Paz,ana@test.com,1\r\n")
        r = self.client.get("/admin/reportes/csv/otro")
        self.assertEqual((r.status_code, r.content), (200, b""))

    def test_csv_sensores(self):
        self.assertEqual(self.client.get("/admin/reportes/csv/sensores").content.decode(),
                         "\ufeffSensor,Tipo,Nivel,Estado,Vehiculo,Cliente\r\nSensor MAF,MAF,10,OK,Onix,Joel\r\n"
                         "Sensor O2,O2,50,Advertencia,Onix,Joel\r\nSensor ECT,ECT,80,Falla,Rio,Ana\r\n"
                         "Prueba,TEST,0,OK,Rio,Ana\r\n")

    def test_csv_estado_null(self):  # Review Focus 1
        Cliente.objects.filter(pk=self.ana.pk).update(estado=None)
        self.assertIn(",ana@test.com,,1\r\n", self.client.get("/admin/reportes/csv/clientes").content.decode())

    def test_csv_escapa_comas_y_comillas(self):  # lote A
        Cliente.objects.filter(pk=self.ana.pk).update(apellido_cliente='Paz, "Jr"')
        filas = list(csv.reader(io.StringIO(self.client.get("/admin/reportes/csv/clientes").content.decode("utf-8-sig"))))
        self.assertEqual(filas[2], [str(self.ana.pk), 'Ana Paz, "Jr"', "ana@test.com", "bloqueado", "1"])

    def test_excel(self):
        r = self.client.get("/admin/reportes/excel?tipoPlaca=publico&nivelMin=")
        self.assertEqual(r["Content-Disposition"], "attachment; filename=reporte-autosen.xlsx")
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Clientes", "Sensores", "Vehículos con fallas", "Clientes bloqueados"])

        def filas(hoja):
            return [list(f) for f in wb[hoja].iter_rows(values_only=True)]

        self.assertEqual(filas("Clientes"), [["ID", "Nombre", "Correo", "Estado", "Vehículos"],
                                             [self.joel.pk, "Joel Moncaleano", "joel@test.com", "activo", 1]])
        self.assertEqual(filas("Sensores")[0], ["Sensor", "Tipo", "Vehículo", "Cliente", "Nivel", "Estado", "Diagnóstico"])
        self.assertEqual(filas("Sensores")[2],
                         ["Sensor O2", "O2", "Onix", "Joel Moncaleano", 50, "ADVERTENCIA", "Falla en flujo de aire"])
        # Con el filtro "publico" solo cuenta Joel: sin vehículos en falla ni clientes bloqueados (revisión final).
        self.assertEqual((len(filas("Sensores")), len(filas("Vehículos con fallas")), len(filas("Clientes bloqueados"))),
                         (3, 1, 1))
        wb = load_workbook(BytesIO(self.client.get("/admin/reportes/excel").content))
        self.assertEqual(filas("Vehículos con fallas")[1], ["Rio", "KIA111", "Ana Paz"])
        self.assertEqual(filas("Clientes bloqueados")[1], ["Ana Paz", "ana@test.com", 1])

    def test_seleccion_del_pdf_respeta_todos_los_filtros(self):  # F5
        params = {"q": "paz", "estadoCliente": None, "marca": "Kia", "tipoPlaca": "particular", "estadoSensor": "falla",
                  "tipoSensor": None, "nivelMin": None, "nivelMax": None}
        self.assertEqual(datos_pdf_general(params), ([self.ana], [self.s3]))
        solo_joel = {**params, "q": "joel", "marca": None, "tipoPlaca": None, "estadoSensor": None}
        self.assertEqual(datos_pdf_general(solo_joel), ([self.joel], [self.s1, self.s2]))  # revisión final

    def test_pdfs(self):
        r = self.client.get("/admin/reportes/pdf?q=&nivelMin=")
        self.assertEqual((r["Content-Type"], r["Content-Disposition"]),
                         ("application/pdf", "attachment; filename=reporte-autosen.pdf"))
        self.assertTrue(r.content.startswith(b"%PDF"))
        r = self.client.get(f"/admin/reportes/cliente/{self.joel.pk}/pdf")
        self.assertEqual(r["Content-Disposition"], "attachment; filename=reporte-joel.pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_pdf_cliente_sin_placa_no_escribe_null(self):  # lote A
        crear_vehiculo(self.joel, placa=None)
        with mock.patch("reportes.exportar._parrafo", wraps=exportar._parrafo) as parrafo:
            self.client.get(f"/admin/reportes/cliente/{self.joel.pk}/pdf")
        textos = [llamada.args[0] for llamada in parrafo.call_args_list]
        self.assertTrue(any(t.startswith("Placa:") for t in textos))
        self.assertFalse(any("null" in t for t in textos), textos)

    def test_pdf_con_texto_unicode(self):  # Review Focus 3
        c = crear_cliente(correo="inigo@test.com", nombre="Íñigo", apellido="Peña")
        crear_vehiculo(c, placa="ÑAN123", nombre="Camión — azul", tipo_placa="Público")
        self.assertTrue(self.client.get(f"/admin/reportes/cliente/{c.pk}/pdf").content.startswith(b"%PDF"))
        self.assertTrue(self.client.get("/admin/reportes/pdf").content.startswith(b"%PDF"))
